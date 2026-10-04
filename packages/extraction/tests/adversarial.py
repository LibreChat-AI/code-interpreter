"""Create bounded adversarial test documents inside the disposable test container."""
from pathlib import Path
import struct
import zipfile
from docx import Document
from pypdf import PdfReader, PdfWriter

root = Path('/tests/fixtures')
(root / 'bad.docx').write_bytes(b'PK\x03\x04not a zip')
with zipfile.ZipFile(root / 'entries.docx', 'w') as file:
    for index in range(513):
        file.writestr(str(index), 'x')
with zipfile.ZipFile(root / 'duplicate.docx', 'w') as file:
    file.writestr('same', 'a')
    file.writestr('same', 'b')
with zipfile.ZipFile(root / 'bomb.docx', 'w', compression=zipfile.ZIP_DEFLATED) as file:
    file.writestr('content.xml', b'x' * (8 * 1024 * 1024 + 1))
with zipfile.ZipFile(root / 'aggregate.docx', 'w', compression=zipfile.ZIP_DEFLATED) as file:
    for index in range(5):
        file.writestr(str(index), b'x' * (7 * 1024 * 1024))
for name, value in [('utf8.docx', b'\xff'), ('entities.docx', b'<!DOCTYPE x [<!ENTITY x "y">]><x/>')]:
    with zipfile.ZipFile(root / name, 'w') as file:
        file.writestr('word/document.xml', value)
with zipfile.ZipFile(root / 'forged.docx', 'w', compression=zipfile.ZIP_DEFLATED) as file:
    file.writestr('content.xml', 'forged-content' * 20)
value = bytearray((root / 'forged.docx').read_bytes())
central = value.index(b'PK\x01\x02')
struct.pack_into('<I', value, central + 24, 1)
(root / 'forged.docx').write_bytes(value)
document = Document()
document.add_paragraph('x' * (1024 * 1024 + 1))
document.save(root / 'output.docx')
document = Document()
document.add_paragraph('x' * (1024 * 1024 - 1))
document.save(root / 'exact.docx')
writer = PdfWriter()
writer.append(root / 'pages.pdf')
writer.encrypt('private-test-password')
writer.write(root / 'encrypted.pdf')
writer = PdfWriter()
for index in range(129):
    writer.add_blank_page(width=10, height=10)
writer.write(root / 'too-many-pages.pdf')
print('Adversarial documents generated')
