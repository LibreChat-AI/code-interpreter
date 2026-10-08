"""Regenerate synthetic, redistributable fixtures with the pinned worker libraries."""
from pathlib import Path
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from docx import Document

root = Path(__file__).parent

def pdf(name, pages, compressed=False, padding=0):
    writer = PdfWriter()
    font = writer._add_object(DictionaryObject({
        NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
        NameObject('/BaseFont'): NameObject('/Helvetica'),
    }))
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(b' ' * padding + b'BT /F1 12 Tf 72 720 Td (' + text.encode() + b') Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream.flate_encode() if compressed else stream)
    with (root / name).open('wb') as file:
        writer.write(file)

pdf('pages.pdf', ['First page text', 'Second page text'])
pdf('compressed.pdf', ['Compressed page text'], True)
pdf('empty.pdf', [''])
pdf('inflate.pdf', [' ' * (8 * 1024 * 1024 + 1)], True)
pdf('aggregate.pdf', ['small'] * 5, True, padding=7 * 1024 * 1024)
document = Document()
document.add_paragraph('First paragraph')
table = document.add_table(rows=1, cols=2)
table.cell(0, 0).text = 'Left cell'
table.cell(0, 1).text = 'Right cell'
document.add_paragraph('Last paragraph')
document.save(root / 'body.docx')
Document().save(root / 'empty.docx')
print('Generated PDF and DOCX fixtures')
